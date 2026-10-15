from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.catalog import importer
from apps.catalog.models import Product
from apps.common.units import Unit, parse_unit

TEMPLATE_URL = "/api/products/import-template/"
PREVIEW_URL = "/api/products/import/preview/"
APPLY_URL = "/api/products/import/apply/"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_bytes(rows, header=("nombre", "unidad", "categoria"), leading=()):
    workbook = Workbook()
    sheet = workbook.active
    for row in leading:
        sheet.append(row)
    if header is not None:
        sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(content, name="productos.xlsx"):
    return SimpleUploadedFile(name, content, content_type=XLSX)


def post_file(client, url, content, name="productos.xlsx"):
    return client.post(url, {"file": upload(content, name)}, format="multipart")


def row_names(data, group):
    return [row["name"] for row in data[group]]


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("unidad", Unit.UNIDAD),
        ("Unidades", Unit.UNIDAD),
        ("UND", Unit.UNIDAD),
        ("und.", Unit.UNIDAD),
        ("lb", Unit.LIBRA),
        ("Libras", Unit.LIBRA),
        ("kg", Unit.KILO),
        ("Kilos", Unit.KILO),
        ("gr", Unit.GRAMOS),
        ("g", Unit.GRAMOS),
        ("Gramo", Unit.GRAMOS),
        ("$", Unit.PESOS),
        ("pesos", Unit.PESOS),
        ("Pesos ($)", Unit.PESOS),
        ("paq", Unit.PAQUETE),
        ("PAQUETES", Unit.PAQUETE),
        ("caja", None),
        ("", None),
    ],
)
def test_parse_unit_is_forgiving(raw, expected):
    assert parse_unit(raw) == expected


@pytest.mark.django_db
class TestTemplate:
    def test_admin_downloads_template(self, client_for, admin_a):
        response = client_for(admin_a).get(TEMPLATE_URL)
        assert response.status_code == 200
        assert response["Content-Type"] == XLSX
        assert "attachment" in response["Content-Disposition"]
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == ("nombre", "unidad", "categoria")
        assert len(rows) == 3
        units_sheet = workbook.worksheets[1]
        listed = [row[0] for row in units_sheet.iter_rows(min_row=2, values_only=True) if row[0]]
        assert set(Unit.values) <= set(listed)

    def test_template_is_importable(self, client_for, admin_a):
        content = client_for(admin_a).get(TEMPLATE_URL).content
        data = post_file(client_for(admin_a), PREVIEW_URL, content).data
        assert data["counts"]["new"] == 2
        assert data["counts"]["invalid"] == 0


@pytest.mark.django_db
class TestPreview:
    def test_classifies_rows(self, client_for, admin_a, branch_a, branch_b):
        Product.objects.create(branch=branch_a, name="Papaya", default_unit=Unit.UNIDAD, category="Frutas")
        Product.objects.create(branch=branch_a, name="Tomate", default_unit=Unit.LIBRA, category="Verduras")
        Product.objects.create(branch=branch_b, name="Cebolla", default_unit=Unit.LIBRA)
        content = xlsx_bytes(
            [
                ("Cebolla", "lbs", "Verduras"),  # new: the other branch's product does not count
                ("PAPAYA", "kg", "Frutas"),  # updated: unit and display name change
                ("tomate", "Libras", ""),  # unchanged except casing -> updated name
                ("Tomate ", "lb", "Verduras"),  # duplicate in file
                ("", "und", "Granos"),  # missing name
                ("Arroz", "caja", ""),  # unknown unit
                (None, None, None),  # blank rows are ignored
            ]
        )
        response = post_file(client_for(admin_a), PREVIEW_URL, content)
        assert response.status_code == 200
        data = response.data
        assert data["counts"] == {"new": 1, "updated": 2, "unchanged": 0, "invalid": 3, "total": 6}
        assert data["new"][0] == {"row": 2, "name": "Cebolla", "unit": "libra", "category": "Verduras"}

        papaya = next(row for row in data["updated"] if row["name"] == "PAPAYA")
        assert {"field": "default_unit", "from": "unidad", "to": "kilo"} in papaya["changes"]
        assert {"field": "name", "from": "Papaya", "to": "PAPAYA"} in papaya["changes"]
        assert papaya["reactivated"] is False

        invalid = {row["row"]: row["errors"] for row in data["invalid"]}
        assert "fila 4" in invalid[5][0]
        assert invalid[6] == ["Falta el nombre del producto."]
        assert "caja" in invalid[7][0]
        # Preview never writes.
        assert Product.objects.filter(branch=branch_a).count() == 2

    def test_unchanged_row(self, client_for, admin_a, branch_a):
        Product.objects.create(branch=branch_a, name="Papaya", default_unit=Unit.UNIDAD, category="Frutas")
        data = post_file(client_for(admin_a), PREVIEW_URL, xlsx_bytes([("Papaya", "und", "")])).data
        assert data["counts"]["unchanged"] == 1
        # A blank category keeps the current one instead of erasing it.
        assert data["unchanged"][0]["category"] == "Frutas"

    def test_blank_unit_defaults_to_unidad_for_new_products(self, client_for, admin_a):
        data = post_file(client_for(admin_a), PREVIEW_URL, xlsx_bytes([("Pan", None, None)])).data
        assert data["new"][0]["unit"] == Unit.UNIDAD

    def test_inactive_product_is_reactivated(self, client_for, admin_a, branch_a):
        Product.objects.create(branch=branch_a, name="Papaya", is_active=False)
        data = post_file(client_for(admin_a), PREVIEW_URL, xlsx_bytes([("Papaya", "unidad", "")])).data
        assert data["counts"]["updated"] == 1
        assert data["updated"][0]["reactivated"] is True

    def test_headers_are_case_and_accent_insensitive_and_can_follow_a_title(self, client_for, admin_a):
        # Columns are matched by header, not position.
        content = xlsx_bytes(
            [("Frutas", "Banano", "Lb")],
            header=("CATEGORÍA", "Nombre ", "UNIDAD"),
            leading=[("Lista de productos",), ()],
        )
        data = post_file(client_for(admin_a), PREVIEW_URL, content).data
        assert data["new"] == [{"row": 4, "name": "Banano", "unit": "libra", "category": "Frutas"}]


@pytest.mark.django_db
class TestFileValidation:
    def test_missing_file(self, client_for, admin_a):
        response = client_for(admin_a).post(PREVIEW_URL, {}, format="multipart")
        assert response.status_code == 400
        assert "file" in response.data

    @pytest.mark.parametrize("name", ["productos.xls", "productos.csv", "productos.pdf"])
    def test_rejects_other_extensions(self, client_for, admin_a, name):
        response = post_file(client_for(admin_a), PREVIEW_URL, xlsx_bytes([("Pan", "", "")]), name)
        assert response.status_code == 400
        assert response.data["code"] == "file_type"

    def test_rejects_corrupt_file(self, client_for, admin_a):
        response = post_file(client_for(admin_a), PREVIEW_URL, b"esto no es un excel")
        assert response.status_code == 400
        assert response.data["code"] == "file_unreadable"

    def test_rejects_large_file(self, client_for, admin_a, monkeypatch):
        monkeypatch.setattr(importer, "MAX_FILE_BYTES", 100)
        response = post_file(client_for(admin_a), PREVIEW_URL, xlsx_bytes([("Pan", "", "")]))
        assert response.status_code == 400
        assert response.data["code"] == "file_too_large"

    def test_rejects_too_many_rows(self, client_for, admin_a, monkeypatch):
        monkeypatch.setattr(importer, "MAX_ROWS", 3)
        content = xlsx_bytes([(f"Producto {i}", "", "") for i in range(4)])
        response = post_file(client_for(admin_a), PREVIEW_URL, content)
        assert response.status_code == 400
        assert response.data["code"] == "too_many_rows"

    def test_rejects_missing_name_header(self, client_for, admin_a):
        content = xlsx_bytes([("Pan", "und")], header=("articulo", "unidad"))
        response = post_file(client_for(admin_a), PREVIEW_URL, content)
        assert response.status_code == 400
        assert response.data["code"] == "missing_header"

    def test_rejects_file_without_rows(self, client_for, admin_a):
        response = post_file(client_for(admin_a), PREVIEW_URL, xlsx_bytes([]))
        assert response.status_code == 400
        assert response.data["code"] == "empty_file"


@pytest.mark.django_db
class TestApply:
    def test_applies_upsert_and_is_idempotent(self, client_for, admin_a, branch_a):
        Product.objects.create(branch=branch_a, name="Papaya", default_unit=Unit.UNIDAD)
        Product.objects.create(branch=branch_a, name="Mango", is_active=False)
        content = xlsx_bytes(
            [
                ("Papaya", "kg", "Frutas"),
                ("Mango", "und", ""),
                ("Leche", "", "Lácteos"),
                ("Arroz", "caja", ""),
            ]
        )
        client = client_for(admin_a)
        response = post_file(client, APPLY_URL, content)
        assert response.status_code == 200
        assert response.data["counts"] == {
            "created": 1, "updated": 2, "reactivated": 1, "unchanged": 0, "invalid": 1,
        }
        assert response.data["invalid"][0]["row"] == 5

        papaya = Product.objects.get(branch=branch_a, normalized_name="papaya")
        assert (papaya.default_unit, papaya.category) == (Unit.KILO, "Frutas")
        assert Product.objects.get(branch=branch_a, normalized_name="mango").is_active is True
        leche = Product.objects.get(branch=branch_a, normalized_name="leche")
        assert (leche.default_unit, leche.category) == (Unit.UNIDAD, "Lácteos")

        again = post_file(client, APPLY_URL, content)
        assert again.data["counts"] == {
            "created": 0, "updated": 0, "reactivated": 0, "unchanged": 3, "invalid": 1,
        }
        assert Product.objects.filter(branch=branch_a).count() == 3

    def test_apply_rejects_file_without_valid_rows(self, client_for, admin_a, branch_a):
        response = post_file(client_for(admin_a), APPLY_URL, xlsx_bytes([("", "und", "x")]))
        assert response.status_code == 400
        assert response.data["code"] == "no_valid_rows"
        assert not Product.objects.filter(branch=branch_a).exists()

    def test_super_admin_imports_into_selected_branch(self, client_for, super_admin, branch_a, branch_b):
        client = client_for(super_admin, branch=branch_b)
        response = post_file(client, APPLY_URL, xlsx_bytes([("Pan", "", "")]))
        assert response.status_code == 200
        assert Product.objects.get(normalized_name="pan").branch == branch_b
        assert not Product.objects.filter(branch=branch_a).exists()

    def test_import_is_isolated_per_branch(self, client_for, admin_a, branch_b):
        other = Product.objects.create(branch=branch_b, name="Papaya", default_unit=Unit.UNIDAD)
        post_file(client_for(admin_a), APPLY_URL, xlsx_bytes([("Papaya", "kg", "")]))
        other.refresh_from_db()
        assert other.default_unit == Unit.UNIDAD
        assert Product.objects.filter(normalized_name="papaya").count() == 2


@pytest.mark.django_db
class TestImportPermissions:
    @pytest.mark.parametrize("url", [PREVIEW_URL, APPLY_URL])
    def test_staff_cannot_import(self, client_for, staff_a, url, branch_a):
        response = post_file(client_for(staff_a), url, xlsx_bytes([("Pan", "", "")]))
        assert response.status_code == 403
        assert not Product.objects.exists()

    def test_staff_cannot_download_template(self, client_for, staff_a):
        assert client_for(staff_a).get(TEMPLATE_URL).status_code == 403

    def test_anonymous_cannot_import(self, client):
        response = client.post(PREVIEW_URL, {"file": upload(xlsx_bytes([("Pan", "", "")]))})
        assert response.status_code == 401
